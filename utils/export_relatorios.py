# utils/export_relatorios.py
#
# Exportação para Excel (.xlsx) e PDF via ReportLab.
# Cada relatório tem seu próprio par de funções: exportar_*_excel / exportar_*_pdf.

from __future__ import annotations

import io
from datetime import datetime
from typing import Any
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage,
)
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils.dataptbr import data_db_para_br
from database.supabase_client import supabase_admin, supabase


# ══════════════════════════════════════════════════════════════════════
# TEMA VISUAL — mesma paleta usada no app (sidebar/topo escuro + azul)
# ══════════════════════════════════════════════════════════════════════

_AZUL_ESCURO   = colors.HexColor("#0F2A44")   # cabeçalho / sidebar do app
_AZUL_PRINCIPAL = colors.HexColor("#1565C0")
_AZUL_CLARO    = colors.HexColor("#E3F2FD")
_CINZA_BORDA   = colors.HexColor("#D6DCE5")
_CINZA_TEXTO   = colors.HexColor("#616161")
_TEXTO         = colors.HexColor("#212121")

# hex sem "#" para uso no openpyxl (que espera "RRGGBB"/"AARRGGBB")
_XL_AZUL_ESCURO = "0F2A44"
_XL_AZUL_CLARO  = "E3F2FD"
_XL_BORDA       = "D6DCE5"
_XL_TEXTO_CLARO = "FFFFFF"

# Mesma logo já usada no app (app/layout.py)
_LOGO_BUCKET = "Heringer"
_LOGO_PATH   = "LogoDocsFlow2-removebg-preview.png"

_logo_bytes_cache: bytes | None = None
_logo_falhou = False


def _obter_logo_bytes() -> bytes | None:
    """
    Baixa a logo do Storage privado (mesmo bucket/arquivo usado no app)
    e guarda em memória — evita baixar de novo a cada exportação.
    Se falhar (bucket/arquivo não encontrado, sem permissão etc.),
    os relatórios continuam sendo gerados normalmente, só sem a logo.
    """
    global _logo_bytes_cache, _logo_falhou

    if _logo_bytes_cache is not None:
        return _logo_bytes_cache
    if _logo_falhou:
        return None

    try:
        client = supabase_admin or supabase
        _logo_bytes_cache = client.storage.from_(_LOGO_BUCKET).download(_LOGO_PATH)
        return _logo_bytes_cache
    except Exception as e:
        print(f"⚠️ Não foi possível carregar a logo para os relatórios: {e}")
        _logo_falhou = True
        return None


# ══════════════════════════════════════════════════════════════════════
# HELPERS INTERNOS
# ══════════════════════════════════════════════════════════════════════

def _buffer_excel(df: pd.DataFrame, sheet_name: str, titulo: str) -> io.BytesIO:
    """Gera BytesIO de um .xlsx com logo, título e cabeçalho estilizados."""
    buf = io.BytesIO()

    LINHA_LOGO   = 1
    LINHA_TITULO = 1
    LINHA_GERADO = 2
    LINHA_HEADER = 4   # 0-indexed p/ pandas = 3

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=LINHA_HEADER - 1)
        ws = writer.sheets[sheet_name]
        n_cols = max(len(df.columns), 1)
        ultima_col = get_column_letter(n_cols)

        # ── logo ──
        logo_bytes = _obter_logo_bytes()
        if logo_bytes:
            try:
                img = XLImage(io.BytesIO(logo_bytes))
                largura_original, altura_original = img.width, img.height
                altura_alvo = 40
                img.height = altura_alvo
                img.width = (
                    int(largura_original * (altura_alvo / altura_original))
                    if altura_original else 120
                )
                ws.add_image(img, "A1")
            except Exception as e:
                print(f"⚠️ Logo não pôde ser inserida no Excel: {e}")

        # ── título + data de geração (colunas ao lado da logo) ──
        col_titulo = get_column_letter(min(3, n_cols))
        ws.merge_cells(f"{col_titulo}{LINHA_TITULO}:{ultima_col}{LINHA_TITULO}")
        cel_titulo = ws[f"{col_titulo}{LINHA_TITULO}"]
        cel_titulo.value = titulo
        cel_titulo.font = Font(size=14, bold=True, color=_XL_AZUL_ESCURO)
        cel_titulo.alignment = Alignment(vertical="center")

        ws.merge_cells(f"{col_titulo}{LINHA_GERADO}:{ultima_col}{LINHA_GERADO}")
        cel_gerado = ws[f"{col_titulo}{LINHA_GERADO}"]
        cel_gerado.value = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        cel_gerado.font = Font(size=9, italic=True, color="757575")

        ws.row_dimensions[1].height = 30

        # ── cabeçalho da tabela ──
        preenchimento_header = PatternFill("solid", fgColor=_XL_AZUL_ESCURO)
        fonte_header = Font(bold=True, color=_XL_TEXTO_CLARO)
        borda_fina = Border(*(Side(style="thin", color=_XL_BORDA) for _ in range(4)))

        for col_idx in range(1, n_cols + 1):
            cel = ws.cell(row=LINHA_HEADER, column=col_idx)
            cel.fill = preenchimento_header
            cel.font = fonte_header
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda_fina

        # ── linhas de dados: borda + faixa alternada ──
        preenchimento_alt = PatternFill("solid", fgColor=_XL_AZUL_CLARO)
        primeira_linha_dados = LINHA_HEADER + 1
        ultima_linha_dados   = LINHA_HEADER + len(df)

        for row_idx in range(primeira_linha_dados, ultima_linha_dados + 1):
            impar = (row_idx - primeira_linha_dados) % 2 == 1
            for col_idx in range(1, n_cols + 1):
                cel = ws.cell(row=row_idx, column=col_idx)
                cel.border = borda_fina
                cel.alignment = Alignment(vertical="center", wrap_text=True)
                if impar:
                    cel.fill = preenchimento_alt

        # ── largura automática das colunas ──
        for col_idx in range(1, n_cols + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                [len(str(ws.cell(row=r, column=col_idx).value or ""))
                 for r in range(LINHA_HEADER, ultima_linha_dados + 1)],
                default=8,
            )
            ws.column_dimensions[col_letter].width = min(max_len + 4, 55)

        ws.freeze_panes = f"A{primeira_linha_dados}"

    buf.seek(0)
    return buf


def _rodape(canvas, doc):
    """Rodapé com marca DocsFlow e número de página, em todas as páginas."""
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(_CINZA_TEXTO)
    largura_pagina = doc.pagesize[0]
    canvas.drawString(1.5 * cm, 1 * cm, "DocsFlow System")
    canvas.drawRightString(
        largura_pagina - 1.5 * cm, 1 * cm, f"Página {doc.page}"
    )
    canvas.setStrokeColor(_CINZA_BORDA)
    canvas.line(1.5 * cm, 1.35 * cm, largura_pagina - 1.5 * cm, 1.35 * cm)
    canvas.restoreState()


def _build_pdf(
    titulo: str,
    colunas: list[str],
    linhas: list[list[str]],
    orientacao: str = "retrato",        # "retrato" | "paisagem"
    col_widths: list[float] | None = None,   # larguras em cm (opcional)
) -> io.BytesIO:
    """Monta PDF estilizado com logo, cabeçalho, tabela e rodapé; retorna BytesIO."""
    buf  = io.BytesIO()
    page = landscape(A4) if orientacao == "paisagem" else A4

    doc = SimpleDocTemplate(
        buf, pagesize=page,
        leftMargin=1.5*cm, rightMargin=1.5*cm,
        topMargin=1.5*cm,  bottomMargin=2*cm,
    )

    styles = getSampleStyleSheet()
    style_celula = ParagraphStyle(
        "CELULA",
        parent=styles["BodyText"],
        fontSize=7,
        leading=9,
        wordWrap="CJK",
    )

    story: list[Any] = []

    # ── cabeçalho: logo + título/data lado a lado ──
    logo_bytes = _obter_logo_bytes()
    titulo_bloco = [
        Paragraph(titulo, ParagraphStyle(
            "H", parent=styles["Heading1"],
            fontSize=15, textColor=_AZUL_ESCURO, spaceAfter=2,
        )),
        Paragraph(
            f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            ParagraphStyle("S", parent=styles["Normal"],
                           fontSize=8, textColor=_CINZA_TEXTO),
        ),
    ]

    if logo_bytes:
        try:
            logo_img = RLImage(io.BytesIO(logo_bytes))
            LARGURA_LOGO_MAX = 2.6 * cm
            ALTURA_LOGO_MAX  = 1.3 * cm
            proporcao = (
                logo_img.imageWidth / logo_img.imageHeight
                if logo_img.imageHeight else 1
            )
            if proporcao > (LARGURA_LOGO_MAX / ALTURA_LOGO_MAX):
                # imagem "larga" — a largura máxima manda
                logo_img.drawWidth  = LARGURA_LOGO_MAX
                logo_img.drawHeight = LARGURA_LOGO_MAX / proporcao
            else:
                # imagem "alta"/quadrada — a altura máxima manda
                logo_img.drawHeight = ALTURA_LOGO_MAX
                logo_img.drawWidth  = ALTURA_LOGO_MAX * proporcao

            header_tbl = Table(
                [[logo_img, titulo_bloco]],
                colWidths=[3.2 * cm, None],
            )
            header_tbl.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN",  (0, 0), (0, 0), "LEFT"),
                ("LEFTPADDING",  (0, 0), (-1, -1), 0),
                ("LEFTPADDING",  (1, 0), (1, 0), 10),
            ]))
            story.append(header_tbl)
        except Exception as e:
            print(f"⚠️ Logo não pôde ser inserida no PDF: {e}")
            story.extend(titulo_bloco)
    else:
        story.extend(titulo_bloco)

    # ── linha divisória sob o cabeçalho ──
    linha_divisoria = Table([[""]], colWidths=[page[0] - 3*cm], rowHeights=[0.05*cm])
    linha_divisoria.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 1, _AZUL_ESCURO),
    ]))
    story.append(Spacer(1, 8))
    story.append(linha_divisoria)
    story.append(Spacer(1, 12))

    avail_w = page[0] - 3*cm
    n_cols  = len(colunas)
    if col_widths:
        cw = [w * cm for w in col_widths]
    else:
        cw = [avail_w / n_cols] * n_cols

    table_data = [colunas]

    for linha in linhas:
        nova_linha = []

        for valor in linha:
            texto = str(valor).replace("\n", "<br/>")
            nova_linha.append(Paragraph(texto, style_celula))

        table_data.append(nova_linha)
    tbl = Table(table_data, colWidths=cw, repeatRows=1)
    tbl.setStyle(TableStyle([
        # cabeçalho
        ("BACKGROUND",    (0, 0), (-1, 0),  _AZUL_ESCURO),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("FONTNAME",      (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0),  8),
        ("ALIGN",         (0, 0), (-1, 0),  "CENTER"),
        ("TOPPADDING",    (0, 0), (-1, 0),  7),
        ("BOTTOMPADDING", (0, 0), (-1, 0),  7),
        # dados
        ("FONTNAME",      (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",      (0, 1), (-1, -1), 7),
        ("TEXTCOLOR",     (0, 1), (-1, -1), _TEXTO),
        ("ALIGN",         (0, 1), (-1, -1), "LEFT"),
        ("TOPPADDING",    (0, 1), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
        ("GRID",          (0, 0), (-1, -1), 0.25, _CINZA_BORDA),
        ("LINEBELOW",     (0, 0), (-1, 0),  1,    _AZUL_ESCURO),
        ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, _AZUL_CLARO]),
    ]))

    story.append(tbl)
    doc.build(story, onFirstPage=_rodape, onLaterPages=_rodape)
    buf.seek(0)
    return buf


def _fmt_datas(d: dict, chaves: list[str], campos_data: set[str]) -> list[str]:
    row = []
    for k in chaves:
        v = d.get(k, "") or ""
        if k in campos_data:
            v = data_db_para_br(v)
        row.append(str(v))
    return row


# ══════════════════════════════════════════════════════════════════════
# 1. CLIENTES
# ══════════════════════════════════════════════════════════════════════

_CHAVES_CLI = ["id", "nome", "tipo", "documento", "sigla"]
_COLS_CLI   = ["ID", "Nome", "Tipo", "Documento", "Sigla"]


def exportar_clientes_excel(dados: list[dict]) -> tuple[io.BytesIO, str]:
    df = pd.DataFrame(dados).reindex(columns=_CHAVES_CLI)
    df.columns = _COLS_CLI
    nome = f"relatorio_clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _buffer_excel(df, "Clientes", "Relatório de Clientes"), nome


def exportar_clientes_pdf(dados: list[dict]) -> tuple[io.BytesIO, str]:
    linhas = [_fmt_datas(d, _CHAVES_CLI, set()) for d in dados]
    buf  = _build_pdf("Relatório de Clientes", _COLS_CLI, linhas)
    nome = f"relatorio_clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return buf, nome


# ══════════════════════════════════════════════════════════════════════
# 2. ALERTAS / NOTIFICAÇÕES
# ══════════════════════════════════════════════════════════════════════

_CHAVES_ALT = ["cliente", "contrato", "observacao", "inicio", "vencimento", "dias_antes", "status"]
_COLS_ALT   = ["Cliente", "Contrato", "Observação", "Início", "Vencimento", "Dias antes", "Status"]
_DATAS_ALT  = {"inicio", "vencimento"}


def exportar_alertas_excel(dados: list[dict], status: str = "todos") -> tuple[io.BytesIO, str]:
    df = pd.DataFrame(dados).reindex(columns=_CHAVES_ALT)
    for col in _DATAS_ALT:
        if col in df.columns:
            df[col] = df[col].apply(data_db_para_br)
    df.columns = _COLS_ALT
    nome = f"relatorio_alertas_{status}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _buffer_excel(df, "Alertas", "Relatório de Alertas"), nome


def exportar_alertas_pdf(dados: list[dict], status: str = "todos") -> tuple[io.BytesIO, str]:
    linhas = [_fmt_datas(d, _CHAVES_ALT, _DATAS_ALT) for d in dados]
    buf  = _build_pdf("Relatório de Alertas", _COLS_ALT, linhas, "paisagem")
    nome = f"relatorio_alertas_{status}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return buf, nome


# ══════════════════════════════════════════════════════════════════════
# 3. CONTRATOS  — schema real (nome, indice, data_inicial,
#                               data_assinatura, termo_final)
# ══════════════════════════════════════════════════════════════════════

_CHAVES_CON = [
    "cliente", "nome", "indice",
    "data_inicial", "data_assinatura", "termo_final",
    "tipo_contrato", "valor", "situacao",
]
_COLS_CON = [
    "Cliente", "Nome do Contrato", "Identificador",
    "Data Inicial", "Data Assinatura", "Termo Final",
    "Tipo", "Valor", "Situação",
]
_DATAS_CON = {"data_inicial", "data_assinatura"}
# termo_final é text no schema (pode conter texto ou data) → não converte


def exportar_contratos_excel(dados: list[dict]) -> tuple[io.BytesIO, str]:
    df = pd.DataFrame(dados).reindex(columns=_CHAVES_CON)
    for col in _DATAS_CON:
        if col in df.columns:
            df[col] = df[col].apply(data_db_para_br)
    df.columns = _COLS_CON
    nome = f"relatorio_contratos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _buffer_excel(df, "Contratos", "Relatório de Contratos"), nome


def exportar_contratos_pdf(dados: list[dict]) -> tuple[io.BytesIO, str]:
    linhas = [_fmt_datas(d, _CHAVES_CON, _DATAS_CON) for d in dados]
    # larguras proporcionais: cliente e nome mais largos
    col_w = [3.5, 4.0, 2.0, 2.2, 2.2, 2.5, 2.2, 2.0, 1.8]
    buf  = _build_pdf("Relatório de Contratos", _COLS_CON, linhas,
                      "paisagem", col_widths=col_w)
    nome = f"relatorio_contratos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return buf, nome


# ══════════════════════════════════════════════════════════════════════
# 4. PRAZOS
# ══════════════════════════════════════════════════════════════════════

_CHAVES_PRA = [
    "id", "cliente", "contrato", "tipo_contrato", "tipo", "observacao",
    "meses", "data_criacao", "data_vencimento",
]
_COLS_PRA = [
    "Identificador", "Cliente", "Contrato", "Tipo de Contrato", "Tipo de Prazo", "Observação",
    "Meses", "Início", "Vencimento",
]
_DATAS_PRA = {"data_criacao", "data_vencimento"}


def exportar_prazos_excel(dados: list[dict]) -> tuple[io.BytesIO, str]:
    df = pd.DataFrame(dados).reindex(columns=_CHAVES_PRA)
    for col in _DATAS_PRA:
        if col in df.columns:
            df[col] = df[col].apply(data_db_para_br)
    df.columns = _COLS_PRA
    nome = f"relatorio_prazos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _buffer_excel(df, "Prazos", "Relatório de Prazos"), nome


def exportar_prazos_pdf(dados: list[dict]) -> tuple[io.BytesIO, str]:
    linhas = [_fmt_datas(d, _CHAVES_PRA, _DATAS_PRA) for d in dados]
    buf  = _build_pdf("Relatório de Prazos", _COLS_PRA, linhas, "paisagem")
    nome = f"relatorio_prazos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return buf, nome